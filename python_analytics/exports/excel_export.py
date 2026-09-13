"""Excel export — generates Revenue_Report.xlsx with 6 sheets.

Sheets:
  Summary     — overall revenue, patients, orders, tests overview
  Patients    — monthly patient registrations
  Orders      — daily revenue breakdown
  Payments    — payment history summary
  Doctor Revenue — doctor-wise revenue ranking
  Top Tests   — most-ordered tests
"""

from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
from openpyxl.utils import get_column_letter

from analytics.revenue_analysis import (
    get_daily_revenue,
    get_monthly_revenue,
    get_overall_summary,
)
from analytics.patient_analysis import (
    get_monthly_registrations,
    get_patient_stats,
)
from analytics.test_analysis import get_all_time_top_tests, get_total_tests_count, get_total_orders_count
from analytics.doctor_analysis import get_all_time_doctor_revenue, get_total_doctors_count

# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F2937")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="6B7280")

DATA_FONT = Font(name="Calibri", size=10)
DATA_ALIGNMENT = Alignment(vertical="center")
CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
CURRENCY_ALIGNMENT = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)

ALT_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

CURRENCY_FORMAT = '#,##0.00'
INTEGER_FORMAT = '#,##0'


def _style_header_row(ws: openpyxl.worksheet.worksheet.Worksheet, headers: list[str], row: int = 1) -> None:
    """Write and style a header row."""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _style_data_cell(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    row: int,
    col: int,
    value,
    *,
    is_currency: bool = False,
    is_integer: bool = False,
    center: bool = False,
    alt: bool = False,
) -> None:
    """Write and style a data cell."""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = DATA_FONT
    cell.border = THIN_BORDER
    if is_currency:
        cell.number_format = CURRENCY_FORMAT
        cell.alignment = CURRENCY_ALIGNMENT
    elif is_integer:
        cell.number_format = INTEGER_FORMAT
        cell.alignment = CENTER_ALIGNMENT
    elif center:
        cell.alignment = CENTER_ALIGNMENT
    else:
        cell.alignment = DATA_ALIGNMENT
    if alt:
        cell.fill = ALT_FILL


def _auto_width(ws: openpyxl.worksheet.worksheet.Worksheet, min_width: int = 10, max_width: int = 40) -> None:
    """Auto-fit column widths."""
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 2, max_width))


def _add_title_block(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    title: str,
    row: int = 1,
) -> int:
    """Add a title row and return the next available row."""
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = TITLE_FONT
    cell = ws.cell(row=row + 1, column=1, value=f"Generated on {date.today().strftime('%d %B %Y')}")
    cell.font = SUBTITLE_FONT
    return row + 3


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _build_summary_sheet(wb: openpyxl.Workbook) -> None:
    """Build the Summary sheet."""
    ws = wb.active
    ws.title = "Summary"

    overall = get_overall_summary()
    patients = get_patient_stats()
    total_tests = get_total_tests_count()
    total_orders = get_total_orders_count()
    total_doctors = get_total_doctors_count()

    next_row = _add_title_block(ws, "Lab Management System — Revenue Report")

    # Summary cards
    summary_data = [
        ("Total Revenue", f"₹ {overall.total_amount:,.2f}"),
        ("Paid Amount", f"₹ {overall.paid_amount:,.2f}"),
        ("Pending Amount", f"₹ {overall.pending_amount:,.2f}"),
        ("Total Orders", f"{overall.order_count:,}"),
        ("Total Patients", f"{patients.total_patients:,}"),
        ("Active Tests", f"{total_tests:,}"),
        ("Total Doctors", f"{total_doctors:,}"),
        ("Avg Revenue / Day", f"₹ {overall.avg_revenue_per_day:,.2f}"),
        ("New Patients (This Month)", f"{patients.new_this_month:,}"),
        ("New Patients (This Year)", f"{patients.new_this_year:,}"),
        ("Patient Growth (YoY)", f"{patients.growth_percent:+.1f}%"),
    ]

    _style_header_row(ws, ["Metric", "Value"], next_row)
    next_row += 1
    for i, (label, value) in enumerate(summary_data):
        alt = i % 2 == 1
        _style_data_cell(ws, next_row, 1, label, alt=alt)
        _style_data_cell(ws, next_row, 2, value, alt=alt)
        next_row += 1

    _auto_width(ws)


def _build_patients_sheet(wb: openpyxl.Workbook) -> None:
    """Build the Patients sheet."""
    ws = wb.create_sheet("Patients")
    next_row = _add_title_block(ws, "Monthly Patient Registrations")

    monthly = get_monthly_registrations(months=12)

    _style_header_row(ws, ["Month", "New Patients", "Cumulative Total"], next_row)
    next_row += 1

    for i, entry in enumerate(monthly):
        alt = i % 2 == 1
        _style_data_cell(ws, next_row, 1, entry.month, alt=alt)
        _style_data_cell(ws, next_row, 2, entry.new_patients, is_integer=True, alt=alt)
        _style_data_cell(ws, next_row, 3, entry.cumulative_total, is_integer=True, alt=alt)
        next_row += 1

    _auto_width(ws)


def _build_orders_sheet(wb: openpyxl.Workbook) -> None:
    """Build the Orders sheet with daily revenue."""
    ws = wb.create_sheet("Orders")
    next_row = _add_title_block(ws, "Daily Revenue Breakdown")

    daily = get_daily_revenue()

    _style_header_row(
        ws,
        ["Date", "Orders", "Total (₹)", "Paid (₹)", "Pending (₹)"],
        next_row,
    )
    next_row += 1

    for i, entry in enumerate(daily):
        alt = i % 2 == 1
        _style_data_cell(ws, next_row, 1, entry.date, alt=alt)
        _style_data_cell(ws, next_row, 2, entry.order_count, is_integer=True, alt=alt)
        _style_data_cell(ws, next_row, 3, entry.total_amount, is_currency=True, alt=alt)
        _style_data_cell(ws, next_row, 4, entry.paid_amount, is_currency=True, alt=alt)
        _style_data_cell(ws, next_row, 5, entry.pending_amount, is_currency=True, alt=alt)
        next_row += 1

    # Totals row
    if daily:
        total_orders = sum(e.order_count for e in daily)
        total_rev = sum(e.total_amount for e in daily)
        total_paid = sum(e.paid_amount for e in daily)
        total_pending = sum(e.pending_amount for e in daily)

        _style_data_cell(ws, next_row, 1, "TOTAL", alt=True)
        _style_data_cell(ws, next_row, 2, total_orders, is_integer=True, alt=True)
        _style_data_cell(ws, next_row, 3, total_rev, is_currency=True, alt=True)
        _style_data_cell(ws, next_row, 4, total_paid, is_currency=True, alt=True)
        _style_data_cell(ws, next_row, 5, total_pending, is_currency=True, alt=True)

    _auto_width(ws)


def _build_payments_sheet(wb: openpyxl.Workbook) -> None:
    """Build the Payments sheet with monthly summary."""
    ws = wb.create_sheet("Payments")
    next_row = _add_title_block(ws, "Monthly Payment Summary")

    monthly = get_monthly_revenue(months=12)

    _style_header_row(
        ws,
        ["Month", "Orders", "Total (₹)", "Paid (₹)", "Pending (₹)", "Collection %"],
        next_row,
    )
    next_row += 1

    for i, entry in enumerate(monthly):
        alt = i % 2 == 1
        collection_pct = (
            round((entry.paid_amount / entry.total_amount) * 100, 1)
            if entry.total_amount > 0
            else 0.0
        )
        _style_data_cell(ws, next_row, 1, entry.period, alt=alt)
        _style_data_cell(ws, next_row, 2, entry.order_count, is_integer=True, alt=alt)
        _style_data_cell(ws, next_row, 3, entry.total_amount, is_currency=True, alt=alt)
        _style_data_cell(ws, next_row, 4, entry.paid_amount, is_currency=True, alt=alt)
        _style_data_cell(ws, next_row, 5, entry.pending_amount, is_currency=True, alt=alt)
        _style_data_cell(ws, next_row, 6, f"{collection_pct}%", center=True, alt=alt)
        next_row += 1

    _auto_width(ws)


def _build_doctor_revenue_sheet(wb: openpyxl.Workbook) -> None:
    """Build the Doctor Revenue sheet."""
    ws = wb.create_sheet("Doctor Revenue")
    next_row = _add_title_block(ws, "Doctor-wise Revenue (All Time)")

    doctors = get_all_time_doctor_revenue(limit=50)

    _style_header_row(
        ws,
        ["Doctor", "Patients", "Orders", "Total (₹)", "Paid (₹)", "Pending (₹)"],
        next_row,
    )
    next_row += 1

    for i, doc in enumerate(doctors):
        alt = i % 2 == 1
        _style_data_cell(ws, next_row, 1, doc.doctor_name, alt=alt)
        _style_data_cell(ws, next_row, 2, doc.patient_count, is_integer=True, alt=alt)
        _style_data_cell(ws, next_row, 3, doc.order_count, is_integer=True, alt=alt)
        _style_data_cell(ws, next_row, 4, doc.total_revenue, is_currency=True, alt=alt)
        _style_data_cell(ws, next_row, 5, doc.paid_revenue, is_currency=True, alt=alt)
        _style_data_cell(ws, next_row, 6, doc.pending_revenue, is_currency=True, alt=alt)
        next_row += 1

    _auto_width(ws)


def _build_top_tests_sheet(wb: openpyxl.Workbook) -> None:
    """Build the Top Tests sheet."""
    ws = wb.create_sheet("Top Tests")
    next_row = _add_title_block(ws, "Top Tests by Order Count (All Time)")

    tests = get_all_time_top_tests(limit=20)

    _style_header_row(
        ws,
        ["Test Name", "Orders", "Revenue (₹)", "Status"],
        next_row,
    )
    next_row += 1

    for i, test in enumerate(tests):
        alt = i % 2 == 1
        _style_data_cell(ws, next_row, 1, test.test_name, alt=alt)
        _style_data_cell(ws, next_row, 2, test.order_count, is_integer=True, alt=alt)
        _style_data_cell(ws, next_row, 3, test.total_revenue, is_currency=True, alt=alt)
        _style_data_cell(ws, next_row, 4, "Active" if test.is_active else "Inactive", center=True, alt=alt)
        next_row += 1

    _auto_width(ws)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate_excel_report(output_dir: Optional[str] = None) -> str:
    """Generate the full Revenue_Report.xlsx and return the file path.

    Args:
        output_dir: Directory to write the file.  Defaults to python_analytics/output/.

    Returns:
        Absolute path to the generated .xlsx file.

    Raises:
        IOError: If the file cannot be written.
    """
    if output_dir is None:
        output_dir = str(Path(__file__).resolve().parent.parent / "output")

    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "Revenue_Report.xlsx")

    wb = openpyxl.Workbook()

    _build_summary_sheet(wb)
    _build_patients_sheet(wb)
    _build_orders_sheet(wb)
    _build_payments_sheet(wb)
    _build_doctor_revenue_sheet(wb)
    _build_top_tests_sheet(wb)

    wb.save(output_path)

    return os.path.abspath(output_path)